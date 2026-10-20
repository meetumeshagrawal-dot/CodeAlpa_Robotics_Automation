import docx, io, copy
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/root/.claude/uploads/4563bebe-9fc3-55e1-9c21-f6f0e90f1f55/c954394c-Task3_Conveyor_Automation_Report.docx"
BASE = "/home/claude/Task3_Conveyor_v2/"
d = docx.Document(SRC)
# logo = first inline image of the sample
logo = d.inline_shapes[0]._inline.graphic.graphicData.pic.blipFill.blip.embed
logo_bytes = d.part.related_parts[logo].blob
body = d.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"): body.remove(el)
sec = d.sections[0]
sec.left_margin = sec.right_margin = Inches(0.75); sec.top_margin = Inches(0.6); sec.bottom_margin = Inches(0.65)
d.styles["Normal"].font.size = Pt(10)
d.styles["Normal"].paragraph_format.space_after = Pt(3)

def P(text="", bold=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, italic=False, after=None):
    p = d.add_paragraph(); p.alignment = align
    if bold: r = p.add_run(bold); r.bold = True; r.font.size = Pt(size) if size else None
    r = p.add_run(text); r.italic = italic
    if size: r.font.size = Pt(size)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p
def H(t):
    h = d.add_heading(t, level=1); h.paragraph_format.space_before = Pt(7); h.paragraph_format.space_after = Pt(2)
    for r in h.runs: r.font.size = Pt(12.5)
def shade(c, col):
    tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), col); tcPr.append(sh)
def table(rows, widths, size=8.5, head=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, x in enumerate(row):
            c = t.cell(i, j); c.width = Inches(widths[j]); c.text = ""; pp = c.paragraphs[0]; pp.paragraph_format.space_after = Pt(0)
            r = pp.add_run(str(x)); r.font.size = Pt(size)
            if head and i == 0: r.bold = True; shade(c, "DCE6F2")
    t.autofit = False
    for j, w in enumerate(widths): t.columns[j].width = Inches(w)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    sp = d.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = Pt(4)
    return t
def img(par, path, w):
    par.add_run().add_picture(path, width=Inches(w))
def caption(t, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = d.add_paragraph(); p.alignment = align; r = p.add_run(t); r.italic = True; r.font.size = Pt(8.5); p.paragraph_format.space_after = Pt(4)

# ---------------- title block
t = d.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False; t.columns[0].width = Inches(1.2); t.columns[1].width = Inches(5.8)
c0, c1 = t.cell(0, 0), t.cell(0, 1); c0.width = Inches(1.2); c1.width = Inches(5.8)
c0.paragraphs[0].add_run().add_picture(io.BytesIO(logo_bytes), width=Inches(0.95))
p = c1.paragraphs[0]; r = p.add_run("CODEALPHA ROBOTICS & AUTOMATION INTERNSHIP"); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = docx.shared.RGBColor(0x1F, 0x4E, 0x79)
p = c1.add_paragraph(); r = p.add_run("TASK 3: AUTOMATION SYSTEM DESIGN"); r.bold = True; r.font.size = Pt(16)
p = c1.add_paragraph(); r = p.add_run("Sensor-Based Conveyor Belt Automation: circuit, flowchart and tested control logic"); r.italic = True; r.font.size = Pt(10.5)
p = c1.add_paragraph(); r = p.add_run("Meet Agrawal  |  B.Tech. Robotics & Artificial Intelligence  |  Ramdeobaba University, Nagpur  |  Roll No. 17  |  2026–27"); r.font.size = Pt(8.5)
for pp in c1.paragraphs: pp.paragraph_format.space_after = Pt(1)

H("Aim")
P("To design a conveyor that stops when an object reaches a work station, holds it for a processing step (count, inspect or sort), then resumes, with a working emergency stop and fault handling. The deliverables are a block diagram, a wiring (circuit) diagram, a flowchart and state diagram, and an Arduino sketch whose logic was compiled and tested on a simulated belt.")

H("System Overview and Components")
P("A 12 V DC gear motor drives the belt. A photoelectric sensor (IR1) at the station detects objects; a second sensor (IR2) at the exit counts delivered items. An Arduino Uno reads the sensors and a START button, and switches the motor through a logic-level MOSFET. A two-contact emergency stop breaks motor power directly in hardware and also reports its state to the Arduino, so safety does not depend on software alone.")
table([["Part", "Value / type", "Function"],
       ["Supply, F1", "12 V DC 2 A, 2 A fuse", "Powers motor; fuse protects wiring"],
       ["S1 E-stop", "Mushroom switch, 2 NC contacts", "S1a cuts motor power; S1b tells the Arduino (open or broken wire = stop)"],
       ["Buck converter", "LM2596 set to 5 V", "Powers Arduino and sensors from the fused 12 V"],
       ["Controller", "Arduino Uno (a PLC in industry)", "Runs the state machine"],
       ["Q1, R1, R2, D1", "IRLZ44N, 220 Ω, 10 kΩ, 1N5819", "Low-side motor switch; gate resistor; pull-down; flyback diode across motor"],
       ["IR1, IR2", "E18-D80NK, NPN normally open", "Object detection and exit counting; output goes LOW on detection"],
       ["S2", "Push button, normally open", "START / RESET (deliberate restart)"],
       ["LED1/2, BZ1", "Green, red (330 Ω); 5 V buzzer via BC547", "Run, stop/fault and alarm indication"]],
      [1.1, 2.3, 3.6])
tb = d.add_table(rows=1, cols=2); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
tb.autofit = False; tb.columns[0].width = Inches(3.45); tb.columns[1].width = Inches(3.55)
a, b = tb.cell(0, 0), tb.cell(0, 1); a.width = Inches(3.45); b.width = Inches(3.55)
pa = a.paragraphs[0]; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER; img(pa, BASE + "Diagrams/Block_Diagram.png", 3.2)
pc = a.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = pc.add_run("Figure 1. System block diagram."); r.italic = True; r.font.size = Pt(8.5)
pb = b.paragraphs[0]; pb.alignment = WD_ALIGN_PARAGRAPH.CENTER; img(pb, BASE + "Diagrams/Flowchart.png", 3.3)
pc = b.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = pc.add_run("Figure 2. Control flowchart."); r.italic = True; r.font.size = Pt(8.5)

H("Wiring Diagram")
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; img(p, BASE + "Diagrams/Circuit_Diagram.png", 6.0)
caption("Figure 3. Wiring diagram of the prototype (pins match the sketch: D2 object sensor, D3 E-stop monitor, D4 exit sensor, D5 START, D6 motor gate, D8/D9 LEDs, D10 buzzer).")
P("The motor is switched by Q1, never driven from a pin. D1 absorbs the voltage spike when the motor switches off. R2 keeps the gate low while the Arduino resets, so the belt cannot start by itself at power-up. Both sensors are wired to inputs with the internal pull-up, so a detection pulls the pin LOW. The E-stop monitor input reads LOW only while the contact is closed, so a pressed button and a broken cable look the same and both stop the belt.")

H("Control Logic")
table([["State", "Motor", "Moves to the next state when"],
       ["IDLE", "off", "START pressed and E-stop healthy → RUN (nothing starts at power-up)"],
       ["RUN", "on", "object sensor active for 30 ms → PROCESS (belt stops, object counted)"],
       ["PROCESS", "off", "2 s dwell finished → CLEARING (time for count / inspect / sort)"],
       ["CLEARING", "on", "beam clear for 300 ms → RUN; not clear after 5 s → FAULT (jam or stuck sensor)"],
       ["FAULT", "off", "START pressed → CLEARING if an object is still held, otherwise RUN; buzzer sounds"],
       ["ESTOP", "off", "E-stop open in any state → ESTOP at once; restart needs E-stop released AND START"]],
      [0.9, 0.6, 5.5])
P("Inputs are debounced for 30 ms so electrical noise cannot trigger a detection, but an E-stop reading acts immediately. An object that is already counted when the belt is stopped by an E-stop or fault is moved out of the beam before normal detection resumes, so it is never counted twice. The full sketch is Design/Conveyor_Control/Conveyor_Control.ino, and the state diagram is in Diagrams/State_Diagram.png.")

H("Testing")
P("The real sketch was compiled on a PC against a small Arduino shim and run on a simulated belt (100 mm/s, 60 mm objects, sensor beam at 500–540 mm). Five scenarios gave 22 of 22 checks passed. One test initially failed and exposed a real bug (an object left in the beam after an E-stop was counted twice); it was fixed and retested.", after=2)
table([["Scenario", "Result"],
       ["Three objects, normal run", "Belt stopped with the object 500–510 mm along (at the sensor), held 2.00 s, then resumed; 3 counted at the station and 3 at the exit"],
       ["Power-up", "Motor stays off in IDLE until START is pressed"],
       ["Sensor noise", "Fifty 10 ms glitches ignored; a sustained 60 ms signal accepted"],
       ["Emergency stop", "Motor off in the same 1 ms control cycle; releasing alone does not restart; START while held is ignored; works in RUN and PROCESS; no double count on restart"],
       ["Jam / stuck sensor", "FAULT after the 2 s dwell plus the 5 s timeout, motor off, buzzer pulsing; START resets once the jam is removed"]],
      [1.5, 5.5])
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; img(p, BASE + "Test/scenario1_timeline.png", 5.2)
caption("Figure 4. Simulated run of the sketch: three objects stop the belt at the station, are held for 2 s, then released; the exit sensor counts them later.")

H("Safety and Limitations")
P("This was tested in simulation, not on hardware, and the circuit has not been built; component values are a starting point and must be checked against the actual motor and sensors. Size the fuse, MOSFET and flyback diode for the motor's stall current and use a contactor or driver module for larger motors. The software emergency stop is a back-up: the real stop is S1a in the motor power path, and an industrial machine needs a safety-rated relay or safety PLC. Set the sensor position from the belt speed and stopping distance, guard the rollers and isolate power before maintenance. The same states map directly onto a PLC ladder program (sensor and button inputs, motor and lamp outputs, timers for dwell and jam timeout).")
H("Conclusion")
P("A sensor, controller and motor driver were combined into a conveyor that stops at a station, processes, resumes, and handles an emergency stop and a jam. The design is documented with a block diagram, wiring diagram, flowchart and state diagram, and its control logic passed all simulated checks, ready for a hardware build and risk assessment.")
d.save(BASE + "Report/Task3_Conveyor_Automation_Report.docx")
